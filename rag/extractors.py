import os
from typing import Dict, List

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from rag.schemas import ExtractionResult


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".py", ".js", ".java", ".json", ".yaml", ".yml",
    ".csv", ".log", ".ini", ".conf", ".xml", ".html", ".css", ".sql",
}


def extract_file_text(file_path: str) -> ExtractionResult:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in TEXT_EXTENSIONS:
            return ExtractionResult(True, "text", _extract_text(file_path))
        if ext == ".docx":
            return ExtractionResult(True, "docx", _extract_docx(file_path))
        if ext == ".xlsx":
            return ExtractionResult(True, "xlsx", _extract_xlsx(file_path))
        if ext == ".pdf":
            return ExtractionResult(True, "pdf", _extract_pdf(file_path))
        if ext == ".pptx":
            return ExtractionResult(True, "pptx", _extract_pptx(file_path))
        return ExtractionResult(False, ext.lstrip("."), "", error="unsupported_extension")
    except Exception as exc:
        return ExtractionResult(False, ext.lstrip("."), "", error=str(exc))


def _extract_text(file_path: str) -> str:
    encodings = ["utf-8", "utf-8-sig", "gbk", "gb2312", "utf-16", "latin-1"]
    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as handle:
                return handle.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("文本解码失败")


def _extract_docx(file_path: str) -> str:
    document = Document(file_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
    return "\n".join(paragraphs)


def _extract_xlsx(file_path: str) -> str:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sections = []
    for sheet in workbook.worksheets:
        sections.append("# Sheet: {0}".format(sheet.title))
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if values:
                sections.append("\t".join(values))
    return "\n".join(sections)


def _extract_pdf(file_path: str) -> str:
    document = fitz.open(file_path)
    pages = []
    for page in document:
        text = page.get_text("text")
        if text and text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def _extract_pptx(file_path: str) -> str:
    presentation = Presentation(file_path)
    lines: List[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        lines.append("# Slide {0}".format(index))
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                lines.append(text.strip())
    return "\n".join(lines)


def get_supported_extensions(config: Dict) -> List[str]:
    raw_extensions = config.get("rag", {}).get("supported_extensions", [])
    return [str(extension).lower() for extension in raw_extensions if str(extension).strip()]
